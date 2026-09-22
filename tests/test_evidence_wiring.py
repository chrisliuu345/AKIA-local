"""判定凭证接线测试（v4.2.0 P4-1/P4-2/P4-3）

覆盖:
1. 分析记录写入凭证（规则版本 / 参数 / 数据指纹 / 软件版本）
2. 追溯码：顺序无关、内容变化即变码；报告页脚携带
3. 复核模式 diff：同数据重跑一致、改参可见、缺凭证不可比
4. Excel 报告落盘后页脚可回读追溯码
"""

import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pandas as pd
import pytest

pytest.importorskip("PySide6")
pytest.importorskip("openpyxl")

from config.settings import APP_VERSION
from core.evidence import (
    credential_trace_code,
    diff_analysis_records,
    fingerprint_dataframe,
    make_credential,
)
from output.base_report_engine import BaseReportEngine
from output.excel_report import ExcelReport


@pytest.fixture
def window():
    from ui.main_window import MainWindow

    win = MainWindow()
    yield win
    win.close()


class TestStoreAnalysisCredential:
    def test_credential_attached_without_data(self, window):
        window._store_analysis(1, {"analysis_type": "spc.imr",
                                   "params": {"subgroup": 5},
                                   "result": {"chart_type": "I-MR"}})
        credential = window.analyses[1]["credential"]
        assert credential["rule_version"] == "spc.imr"
        assert credential["params"] == {"subgroup": 5}
        assert credential["data_fingerprint"] == "no-data"
        assert credential["app_version"] == APP_VERSION
        assert credential["computed_at"]

    def test_credential_fingerprints_imported_data(self, window):
        frame = pd.DataFrame({"x": [1.0, 2.0, 3.0], "y": [4.0, 5.0, 6.0]})
        window.current_data = {"raw_data": {"Sheet1": frame}}
        window._store_analysis(2, {"analysis_type": "capability",
                                   "params": {}, "result": {}})
        credential = window.analyses[2]["credential"]
        assert credential["data_fingerprint"] == fingerprint_dataframe(frame)
        assert credential["data_fingerprint"].startswith("pd-hash-v1:")

    def test_rule_version_falls_back_to_chart_type(self, window):
        window._store_analysis(3, {"chart_type": "MEWMA", "result": {}})
        assert window.analyses[3]["credential"]["rule_version"] == "MEWMA"

    def test_rerun_records_diff(self, window):
        payload = {"analysis_type": "spc.imr", "params": {"k": 5},
                   "result": {"chart_type": "I-MR", "summary": {"x_bar": 2.0}}}
        window._store_analysis(4, payload)
        assert "rerun_diff" not in window.analyses[4], "首次分析不应有 diff"

        window._store_analysis(4, payload)
        diff = window.analyses[4]["rerun_diff"]
        assert diff["comparable"] is True
        assert diff["identical"] is True

        changed = {"analysis_type": "spc.imr", "params": {"k": 6},
                   "result": {"chart_type": "I-MR", "summary": {"x_bar": 2.0}}}
        window._store_analysis(4, changed)
        diff = window.analyses[4]["rerun_diff"]
        assert diff["identical"] is False
        assert "params" in diff["credential_changes"]


class TestTraceCode:
    @staticmethod
    def _credential(rule: str, value: int = 1):
        return make_credential(rule_version=rule, params={"k": value}).to_dict()

    def test_order_independent(self):
        a = self._credential("spc", 1)
        b = self._credential("msa", 2)
        assert credential_trace_code([a, b]) == credential_trace_code([b, a])

    def test_changes_with_params(self):
        assert credential_trace_code([self._credential("spc", 1)]) != \
            credential_trace_code([self._credential("spc", 2)])

    def test_changes_with_rule_version(self):
        assert credential_trace_code([self._credential("spc", 1)]) != \
            credential_trace_code([self._credential("capability", 1)])

    def test_empty_without_credentials(self):
        assert credential_trace_code([]) == ""
        assert credential_trace_code([None, {}]) == ""

    def test_length_and_charset(self):
        code = credential_trace_code([self._credential("spc")])
        assert len(code) == 12
        assert all(c in "0123456789ABCDEF" for c in code)


class TestRerunDiff:
    @staticmethod
    def _record(rule: str, params: dict, result: dict):
        return {
            "credential": make_credential(rule_version=rule, params=params).to_dict(),
            "result": result,
        }

    def test_identical_rerun(self):
        old = self._record("spc.imr", {"k": 1}, {"cpk": 1.33})
        new = self._record("spc.imr", {"k": 1}, {"cpk": 1.33})
        diff = diff_analysis_records(old, new)
        assert diff["comparable"] is True
        assert diff["identical"] is True
        assert diff["credential_changes"] == {}
        assert diff["result_changed"] is False

    def test_param_change_is_visible(self):
        old = self._record("spc.imr", {"k": 1}, {"cpk": 1.33})
        new = self._record("spc.imr", {"k": 2}, {"cpk": 1.10})
        diff = diff_analysis_records(old, new)
        assert diff["identical"] is False
        assert "params" in diff["credential_changes"]
        assert diff["credential_changes"]["params"]["old"] == {"k": 1}
        assert diff["credential_changes"]["params"]["new"] == {"k": 2}
        assert diff["result_changed"] is True

    def test_missing_credential_is_not_comparable(self):
        diff = diff_analysis_records({"result": {}}, self._record("spc", {}, {}))
        assert diff["comparable"] is False
        assert "缺少凭证" in diff["reason"]


class TestReportTrace:
    def test_footer_contains_trace_code(self):
        footer, _date = BaseReportEngine.footers(trace="ABCD1234EF56")
        assert "追溯码: ABCD1234EF56" in footer
        plain, _ = BaseReportEngine.footers()
        assert "追溯码" not in plain

    def test_credential_trace_from_analyses(self):
        analyses = [
            {"credential": make_credential("spc", {"k": 1}).to_dict()},
            {"credential": make_credential("msa", {"k": 2}).to_dict()},
            {"result": {}},
        ]
        code = BaseReportEngine.credential_trace(analyses)
        assert len(code) == 12
        assert code == credential_trace_code(
            [analyses[0]["credential"], analyses[1]["credential"]])

    def test_excel_report_writes_trace_code(self, tmp_path):
        frame = pd.DataFrame({"x": [1.0, 2.0, 3.0]})
        credential = make_credential(rule_version="spc.imr", params={"k": 5},
                                     df=frame).to_dict()
        analyses = [{
            "analysis_type": "spc.imr",
            "params": {"k": 5},
            "result": {"chart_type": "I-MR", "summary": {"x_bar": 2.0}},
            "credential": credential,
        }]
        output = tmp_path / "report.xlsx"
        ExcelReport().generate(analyses, str(output), title="凭证测试")

        from openpyxl import load_workbook

        workbook = load_workbook(output)
        summary = workbook[workbook.sheetnames[0]]
        footer = str(summary.cell(row=2, column=1).value or "")
        assert "追溯码" in footer
        assert credential_trace_code([credential]) in footer
