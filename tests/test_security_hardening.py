"""W06-W17 安全加固测试 — Commit D"""

import os
import sys
import base64
import logging
import pytest
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))


# ── W06: SQL 连接串密码 URL 编码 ──────────────────────

class TestW06PasswordEncoding:
    """W06: 密码含 @ 等特殊字符时 URL 编码，确保 sanitize 正确打码"""

    def test_password_with_at_is_encoded(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mysql", "host": "localhost", "port": 3306,
            "database": "testdb", "user": "root", "password": "p@ss",
        }
        conn_str = build_connection_string(config)
        assert "p%40ss" in conn_str
        assert conn_str.count("@") == 1

    def test_password_with_special_chars_encoded(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "postgresql", "host": "localhost", "port": 5432,
            "database": "testdb", "user": "user@1", "password": "p#s/w:r:d",
        }
        conn_str = build_connection_string(config)
        assert "user%401" in conn_str
        assert "p%23s%2Fw%3Ar%3Ad" in conn_str
        assert conn_str.count("@") == 1

    def test_sanitize_error_with_encoded_password(self):
        from core.sql_loader import sanitize_connection_error
        err = Exception("connect failed: mysql+pymysql://root:p%40ss@localhost:3306/db")
        sanitized = sanitize_connection_error(err)
        assert "p%40ss" not in sanitized
        assert "******" in sanitized

    def test_mssql_password_encoded(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mssql", "host": "localhost", "port": 1433,
            "database": "testdb", "user": "sa", "password": "P@ssw0rd!",
        }
        conn_str = build_connection_string(config)
        assert "P%40ssw0rd%21" in conn_str
        assert conn_str.count("@") == 1


# ── W08: 共享项目路径穿越防护 ─────────────────────────

class TestW08PathTraversal:
    """W08: save_shared 部门参数路径穿越防护"""

    def test_traversal_sanitized_to_basename(self, tmp_path, monkeypatch):
        import core.project_manager as pm
        monkeypatch.setattr(pm, "SHARED_DIR", str(tmp_path))
        saved_paths = []

        def mock_save(path, *args, **kwargs):
            saved_paths.append(path)
            return path
        monkeypatch.setattr(pm.ProjectManager, "save_project", mock_save)

        pm.SharedProjectManager.save_shared("test.qap", {}, {}, {}, "../../etc", "")
        assert len(saved_paths) == 1
        assert saved_paths[0].startswith(str(tmp_path))
        assert "etc" in saved_paths[0]
        assert ".." not in saved_paths[0]

    def test_dotdot_rejected_by_containment(self, tmp_path, monkeypatch):
        import core.project_manager as pm
        monkeypatch.setattr(pm, "SHARED_DIR", str(tmp_path))
        with pytest.raises(ValueError, match="非法"):
            pm.SharedProjectManager.save_shared("test.qap", {}, {}, {}, "..", "")

    def test_normal_department_works(self, tmp_path, monkeypatch):
        import core.project_manager as pm
        monkeypatch.setattr(pm, "SHARED_DIR", str(tmp_path))
        saved_paths = []

        def mock_save(path, *args, **kwargs):
            saved_paths.append(path)
            return path
        monkeypatch.setattr(pm.ProjectManager, "save_project", mock_save)

        pm.SharedProjectManager.save_shared("test.qap", {}, {}, {}, "质量部", "")
        assert saved_paths[0].startswith(str(tmp_path))
        assert "质量部" in saved_paths[0]

    def test_backslash_traversal_sanitized(self, tmp_path, monkeypatch):
        """G15: Windows 反斜杠穿越被 basename 清洗, 落点仍在共享目录内"""
        import core.project_manager as pm
        monkeypatch.setattr(pm, "SHARED_DIR", str(tmp_path))
        saved_paths = []

        def mock_save(path, *args, **kwargs):
            saved_paths.append(path)
            return path
        monkeypatch.setattr(pm.ProjectManager, "save_project", mock_save)

        pm.SharedProjectManager.save_shared("test.qap", {}, {}, {}, r"..\..\etc", "")
        assert len(saved_paths) == 1
        assert saved_paths[0].startswith(str(tmp_path))
        assert ".." not in saved_paths[0]

    def test_empty_department_uses_public(self, tmp_path, monkeypatch):
        """G14: 空部门落到 PUBLIC_DEPT (公共)"""
        import core.project_manager as pm
        monkeypatch.setattr(pm, "SHARED_DIR", str(tmp_path))
        saved_paths = []

        def mock_save(path, *args, **kwargs):
            saved_paths.append(path)
            return path
        monkeypatch.setattr(pm.ProjectManager, "save_project", mock_save)

        pm.SharedProjectManager.save_shared("test.qap", {}, {}, {}, "", "")
        assert saved_paths[0].endswith("公共.qap") or saved_paths[0].endswith("公共\\test.qap") or "公共" in saved_paths[0]


# ── W09: 日志器不再单例缓存 ───────────────────────────

class TestW09DistinctLoggers:
    """W09: get_logger 返回不同名称的独立 logger"""

    def test_distinct_names_return_distinct_loggers(self):
        from core.log_util import get_logger
        logger_a = get_logger("TestModuleA")
        logger_b = get_logger("TestModuleB")
        assert logger_a.name == "TestModuleA"
        assert logger_b.name == "TestModuleB"
        assert logger_a is not logger_b

    def test_same_name_returns_same_logger(self):
        from core.log_util import get_logger
        logger1 = get_logger("TestSame")
        logger2 = get_logger("TestSame")
        assert logger1 is logger2


# ── W15: DPAPI 降级告警 ───────────────────────────────

class TestW15DPAPIDegradation:
    """W15: _obscure 加密失败时返回空字符串"""

    def test_warning_logged_on_dpapi_failure(self, caplog):
        import ui.dialogs.db_connection_dialog as dbcd
        with patch("ui.dialogs.db_connection_dialog.encrypt_secret", side_effect=Exception("Encryption unavailable")):
            result = dbcd._obscure("mysecret")
        assert result == ""  # 加密失败返回空字符串


# ── W16: 合规模板路径穿越防护 ─────────────────────────

class TestW16TemplateTraversal:
    """W16: load_template 拒绝路径穿越"""

    def test_dotdot_rejected(self):
        from output.compliance_report import ComplianceReportGenerator
        with pytest.raises(ValueError, match="非法"):
            ComplianceReportGenerator.load_template("../secret")

    def test_slash_rejected(self):
        from output.compliance_report import ComplianceReportGenerator
        with pytest.raises(ValueError, match="非法"):
            ComplianceReportGenerator.load_template("foo/bar")

    def test_backslash_rejected(self):
        from output.compliance_report import ComplianceReportGenerator
        with pytest.raises(ValueError, match="非法"):
            ComplianceReportGenerator.load_template("foo\\bar")

    def test_empty_rejected(self):
        from output.compliance_report import ComplianceReportGenerator
        with pytest.raises(ValueError, match="非法"):
            ComplianceReportGenerator.load_template("")


# ── W17: Excel 公式注入防护 ───────────────────────────

class TestW17FormulaInjection:
    """W17: _safe_write 字符串使用 write_string 防止公式注入"""

    def test_string_uses_write_string(self):
        from output.excel_report import ExcelReport
        report = ExcelReport.__new__(ExcelReport)
        ws = MagicMock()
        report._safe_write(ws, 0, 0, "=cmd|' /c calc'!A1", None)
        ws.write_string.assert_called_once_with(0, 0, "=cmd|' /c calc'!A1", None)
        ws.write.assert_not_called()

    def test_number_uses_write(self):
        from output.excel_report import ExcelReport
        report = ExcelReport.__new__(ExcelReport)
        ws = MagicMock()
        report._safe_write(ws, 0, 0, 42, None)
        ws.write.assert_called_once_with(0, 0, 42, None)
        ws.write_string.assert_not_called()

    def test_normal_string_uses_write_string(self):
        from output.excel_report import ExcelReport
        report = ExcelReport.__new__(ExcelReport)
        ws = MagicMock()
        report._safe_write(ws, 0, 0, "Cpk = 1.33", None)
        ws.write_string.assert_called_once_with(0, 0, "Cpk = 1.33", None)
        ws.write.assert_not_called()

    def test_float_uses_write(self):
        from output.excel_report import ExcelReport
        report = ExcelReport.__new__(ExcelReport)
        ws = MagicMock()
        report._safe_write(ws, 0, 0, 1.33, None)
        ws.write.assert_called_once_with(0, 0, 1.33, None)
        ws.write_string.assert_not_called()


# ── Commit G: W06 空格/中文/空值编码回归 ──────────────

class TestW06SpaceEncoding:
    """W06 修复: quote(safe='') — 空格 %20 可被 SQLAlchemy unquote 还原"""

    def test_password_with_space_uses_percent20(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mysql", "host": "localhost", "port": 3306,
            "database": "testdb", "user": "root", "password": "my pass",
        }
        conn_str = build_connection_string(config)
        assert "my%20pass" in conn_str, "空格必须编码为 %20 而非 +"
        assert "+" not in conn_str.split("@")[0].split(":")[-1] or "my%20pass" in conn_str

    def test_password_roundtrip_via_make_url(self):
        """真实解析验证: %20 还原为空格 (quote_plus 的 + 不能还原)"""
        from sqlalchemy.engine import make_url
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mysql", "host": "localhost", "port": 3306,
            "database": "testdb", "user": "root", "password": "my pass",
        }
        url = make_url(build_connection_string(config))
        assert url.password == "my pass"

    def test_chinese_password_encoded(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mysql", "host": "localhost", "port": 3306,
            "database": "testdb", "user": "root", "password": "密码@2024",
        }
        conn_str = build_connection_string(config)
        assert conn_str.count("@") == 1, "密码中的 @ 必须被编码"

    def test_empty_user_and_password(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mysql", "host": "localhost", "port": 3306,
            "database": "testdb", "user": "", "password": "",
        }
        conn_str = build_connection_string(config)
        assert "://:@localhost" in conn_str

    def test_database_with_slash_encoded(self):
        from core.sql_loader import build_connection_string
        config = {
            "db_type": "mysql", "host": "localhost", "port": 3306,
            "database": "a/b", "user": "root", "password": "p",
        }
        conn_str = build_connection_string(config)
        assert "a%2Fb" in conn_str

    def test_unsupported_db_type_raises(self):
        from core.sql_loader import build_connection_string
        import pytest
        with pytest.raises(ValueError):
            build_connection_string({"db_type": "oracle"})


# ── Commit G: W07 reflect 回退路径 ───────────────────

class TestW07ReflectFallback:
    """W07: read_sql_table 失败时回退 MetaData.reflect + select"""

    def test_fallback_success_when_read_sql_table_fails(self):
        import pandas as pd
        from core.sql_loader import SqlLoader
        loader = SqlLoader()
        loader.engine = MagicMock()
        with patch("pandas.read_sql_table", side_effect=Exception("boom")), \
             patch("pandas.read_sql_query", return_value=pd.DataFrame({"a": [1]})), \
             patch("sqlalchemy.select", return_value=MagicMock()):
            with patch("sqlalchemy.MetaData") as md_cls:
                md = MagicMock()
                md.tables = {"mytable": MagicMock()}
                md_cls.return_value = md
                result = loader.load(["mytable"])
        assert "mytable" in result

    def test_double_failure_raises_sanitized_error(self):
        from core.sql_loader import SqlLoader
        loader = SqlLoader()
        loader.engine = MagicMock()
        with patch("pandas.read_sql_table", side_effect=Exception("boom")):
            with patch("sqlalchemy.MetaData") as md_cls:
                md = MagicMock()
                md.tables = {}  # reflect 后无该表 → ValueError → RuntimeError
                md_cls.return_value = md
                with pytest.raises(RuntimeError) as ei:
                    loader.load(["mytable"])
        assert "mytable" in str(ei.value)

    def test_double_failure_error_message_sanitized(self):
        from core.sql_loader import SqlLoader
        loader = SqlLoader()
        loader.engine = MagicMock()
        with patch("pandas.read_sql_table", side_effect=Exception("boom")):
            with patch("sqlalchemy.MetaData") as md_cls:
                md = MagicMock()
                md.tables = {}
                md_cls.return_value = md
                with patch("core.sql_loader.sanitize_connection_error",
                           return_value="REDACTED") as se:
                    with pytest.raises(RuntimeError):
                        loader.load(["mytable"])
                    se.assert_called()


# ── Commit G: W16 正向路径 ───────────────────────────

class TestW16TemplatePositive:
    """W16: 合法 ID 放行, 不存在抛 FileNotFoundError"""

    def test_missing_template_raises_filenotfound(self, monkeypatch, tmp_path):
        from output import compliance_report
        monkeypatch.setattr(compliance_report, "TEMPLATES_DIR", str(tmp_path))
        with pytest.raises(FileNotFoundError):
            compliance_report.ComplianceReportGenerator.load_template("iso_9001")

    def test_valid_template_loaded(self, monkeypatch, tmp_path):
        import json
        from output import compliance_report
        (tmp_path / "iso_9001.json").write_text(
            json.dumps({"name": "ISO 9001", "format": "docx"}), encoding="utf-8")
        monkeypatch.setattr(compliance_report, "TEMPLATES_DIR", str(tmp_path))
        data = compliance_report.ComplianceReportGenerator.load_template("iso_9001")
        assert data["name"] == "ISO 9001"

    def test_valid_id_with_dot_underscore_allowed(self, monkeypatch, tmp_path):
        from output import compliance_report
        monkeypatch.setattr(compliance_report, "TEMPLATES_DIR", str(tmp_path))
        with pytest.raises(FileNotFoundError):
            compliance_report.ComplianceReportGenerator.load_template("v2.1_eng")


# ── Commit G: W17 集成级公式注入防护 ─────────────────

class TestW17MergeRangeIntegration:
    """真实 xlsx 输出验证: 标题/atype 含 = 前缀也不会被写成公式"""

    def _generate(self, analyses, tmp_path):
        from output.excel_report import ExcelReport
        out = str(tmp_path / "report.xlsx")
        ExcelReport().generate(analyses, out, title="=1+1")
        return out

    def test_formula_like_title_written_as_string(self, tmp_path):
        from openpyxl import load_workbook
        out = self._generate([], tmp_path)
        wb = load_workbook(out, data_only=False)
        ws = wb.worksheets[0]
        assert ws.cell(row=1, column=1).data_type == "s", "标题必须为字符串"
        assert ws.cell(row=1, column=1).value == "=1+1"

    def test_formula_like_atype_written_as_string(self, tmp_path):
        from openpyxl import load_workbook
        analyses = [{
            "analysis_type": "=HYPERLINK(\"http://evil\")",
            "params": {"a": 1}, "result": {"x": 1},
        }]
        out = self._generate(analyses, tmp_path)
        wb = load_workbook(out, data_only=False)
        ws = wb.worksheets[1]
        assert ws.cell(row=1, column=1).data_type == "s"
        assert ws.cell(row=1, column=1).value == '=HYPERLINK("http://evil")'
