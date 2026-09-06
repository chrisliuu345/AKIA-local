"""
PowerPoint 图表报告生成器单元测试
"""

import matplotlib
matplotlib.use("Agg")

import numpy as np

import pytest
from pptx import Presentation

from output.chart_engine import figure_to_bytes
from output.pptx_report import PptxReport


@pytest.fixture
def fake_analyses():
    return [
        {
            "analysis_type": "SPC: XBar-R",
            "params": {"列": "直径", "子组": 5},
            "result": {"chart_type": "XBar-R", "summary": "过程稳定，无失控点"},
        },
        {
            "analysis_type": "能力分析",
            "params": {"列": "直径"},
            "result": {"chart_type": "能力图", "summary": "Cpk=1.32"},
        },
    ]


def _make_chart_bytes():
    from output.chart_engine import create_histogram
    fig = create_histogram(np.random.normal(0, 1, 200), title="测试图")
    return figure_to_bytes(fig).getvalue()


@pytest.fixture
def charts_bytes():
    return {
        "SPC: XBar-R": _make_chart_bytes(),
        "能力分析": _make_chart_bytes(),
    }


def test_generate_pptx(tmp_path, fake_analyses, charts_bytes):
    out = tmp_path / "deck.pptx"
    PptxReport().generate(fake_analyses, str(out), charts=charts_bytes)
    assert out.exists() and out.stat().st_size > 0

    prs = Presentation(str(out))
    # 封面 + 每分析一页
    assert len(prs.slides) == 3
    slides_text = []
    pic_count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.shape_type == 13:  # PICTURE
                pic_count += 1
            if shape.has_text_frame:
                slides_text.append(shape.text_frame.text)
    all_text = "".join(slides_text)
    assert "SPC: XBar-R" in all_text
    assert "能力分析" in all_text
    assert "1 / 2" in all_text
    assert pic_count == 2


def test_generate_without_charts(tmp_path, fake_analyses):
    out = tmp_path / "nochart.pptx"
    PptxReport().generate(fake_analyses, str(out), charts=None)
    prs = Presentation(str(out))
    assert len(prs.slides) == 3
    all_text = "".join(
        sh.text_frame.text
        for slide in prs.slides for sh in slide.shapes if sh.has_text_frame)
    assert "该分析无图表" in all_text


def test_custom_title(tmp_path, fake_analyses):
    out = tmp_path / "custom.pptx"
    PptxReport().generate(fake_analyses, str(out), title="月度质量图表")
    prs = Presentation(str(out))
    cover_text = "".join(
        sh.text_frame.text
        for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "月度质量图表" in cover_text
