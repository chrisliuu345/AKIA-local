"""
Word 报告生成器 — 基于 python-docx 生成专业质量分析报告
"""

from datetime import datetime
from typing import Dict, Any, List
from io import BytesIO
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from config.settings import APP_NAME, APP_VERSION, CHART_FIGSIZE
from core.i18n import i18n
from output.base_report_engine import BaseReportEngine
from ui.font_config import get_ui_font

def _hdg(num: int, atype: str) -> str:
    """Generate section heading"""
    return f"{num}、{atype}"


class WordReport(BaseReportEngine):
    """生成 Word (.docx) 格式的质量分析报告"""

    def __init__(self, font: str = None, font_size: int = 11):
        super().__init__()
        self.document = Document()
        self.font = font or get_ui_font()
        self.font_size = font_size
        self._setup_styles()

    def _setup_styles(self):
        style = self.document.styles["Normal"]
        style.font.name = self.font
        style.font.size = Pt(self.font_size)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)

        for level in range(1, 4):
            heading_style = self.document.styles[f"Heading {level}"]
            heading_style.font.name = self.font
            heading_style.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)

    def generate(
        self,
        analyses: List[Dict[str, Any]],
        output_path: str,
        title: str = "",
        subtitle: str = "",
        charts: Dict[str, Any] = None,
    ) -> str:
        if not title:
            title = i18n.tr("report.title")
        doc = self.document

        self._add_title_page(title, subtitle)

        self._add_table_of_contents_placeholder()

        doc.add_heading(i18n.tr("report.overview"), level=1)
        overview = doc.add_paragraph()
        overview.add_run(i18n.tr("report.auto_generated")).font.size = Pt(10)
        footer_text, _ = self.footers()
        doc.add_paragraph(footer_text)
        doc.add_paragraph(f"{i18n.tr('report.analysis_modules')}: {len(analyses)}")
        doc.add_paragraph("")

        for i, atype, params, result in self.iter_analyses(analyses):

            doc.add_heading(_hdg(i+1, atype), level=1)

            if params:
                doc.add_heading(i18n.tr("report.analysis_params"), level=2)
                self._add_params_table(params)

            chart_type = result.get("chart_type", atype)

            if chart_type in ("XBar-R", "XBar-S", "I-MR"):
                doc.add_heading(i18n.tr("report.spc_results"), level=2)
                self._add_spc_results(result)

            elif "Capability" in chart_type or "capability" in atype:
                doc.add_heading(i18n.tr("report.capability_results"), level=2)
                self._add_capability_results(result)

            elif "Gage" in chart_type or "gage" in atype or "MSA" in atype:
                doc.add_heading(i18n.tr("report.msa_results"), level=2)
                self._add_msa_results(result)

            elif "t-Test" in chart_type or "t_test" in atype.lower() or "ANOVA" in chart_type:
                doc.add_heading(i18n.tr("report.hypothesis_results"), level=2)
                self._add_hypothesis_results(result)

            elif "DOE" in chart_type or "factorial" in atype.lower():
                doc.add_heading(i18n.tr("report.doe_results"), level=2)
                self._add_doe_results(result)

            if charts and atype in charts:
                doc.add_heading(i18n.tr("report.charts"), level=2)
                chart_data = charts[atype]
                chart_buf = BytesIO(chart_data) if isinstance(chart_data, bytes) else chart_data
                # W80: 限制图片宽度不超过 A4 有效宽度(6.0 英寸)，按比例缩放
                max_w = 6.0
                w = min(CHART_FIGSIZE[0], max_w)
                h = w * CHART_FIGSIZE[1] / CHART_FIGSIZE[0]
                doc.add_picture(chart_buf, width=Inches(w), height=Inches(h))

            doc.add_page_break()

        self._add_appendix()

        doc.save(output_path)
        return output_path

    def _add_title_page(self, title: str, subtitle: str):
        doc = self.document

        for _ in range(4):
            doc.add_paragraph("")

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(title)
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.name = self.font
        run.element.rPr.rFonts.set(qn("w:eastAsia"), self.font)

        if subtitle:
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run2 = p2.add_run(subtitle)
            run2.font.size = Pt(16)
            run2.font.color.rgb = RGBColor(100, 100, 100)

        doc.add_paragraph("")

        info_lines = [
            f"{i18n.tr('report.generated_by')} {APP_NAME} v{APP_VERSION}",
            f"{i18n.tr('report.date')}: {datetime.now().strftime('%Y-%m-%d')}",
        ]
        for line in info_lines:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(line).font.size = Pt(11)

        doc.add_page_break()

    def _add_table_of_contents_placeholder(self):
        doc = self.document
        doc.add_heading(i18n.tr("report.toc"), level=1)
        doc.add_page_break()

    def _add_params_table(self, params: Dict[str, Any]):
        doc = self.document
        table = doc.add_table(rows=len(params) + 1, cols=2, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        table.cell(0, 0).text = i18n.tr("report.param")
        table.cell(0, 1).text = i18n.tr("report.value")
        for i, (k, v) in enumerate(params.items(), 1):
            table.cell(i, 0).text = str(k)
            table.cell(i, 1).text = str(v)

    def _add_spc_results(self, result: Dict):
        doc = self.document
        summary = result.get("summary", {})

        table = doc.add_table(rows=4, cols=2, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        rows_data = [
            (i18n.tr("report.subgroups"), str(result.get("num_subgroups", summary.get("num_subgroups", "")))),
            (i18n.tr("report.subgroup_size"), str(result.get("subgroup_size", ""))),
            (i18n.tr("report.process_mean"), str(summary.get("xbar_bar", summary.get("x_bar", "")))),
            (i18n.tr("report.sigma_hat"), str(summary.get("sigma_hat", ""))),
        ]
        for i, (k, v) in enumerate(rows_data):
            table.cell(i, 0).text = k
            table.cell(i, 1).text = v

        violations_key = "xbar_chart" if "xbar_chart" in result else "i_chart"
        violations = result.get(violations_key, {}).get("violations", [])
        if violations:
            doc.add_heading(i18n.tr("report.violations"), level=3)
            for v in violations[:20]:
                doc.add_paragraph(f"点 {v.get('point', '')}: {v.get('desc', v.get('type', ''))}",
                                  style="List Bullet")

    def _add_capability_results(self, result: Dict):
        doc = self.document
        cap = result.get("capability", {})
        summary = result.get("summary", {})

        table = doc.add_table(rows=7, cols=2, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        rows = [
            (i18n.tr("report.sample_size"), str(summary.get("sample_size", ""))),
            (i18n.tr("report.mean"), str(summary.get("mean", ""))),
            (i18n.tr("report.std_dev"), str(summary.get("std_overall", ""))),
            ("Cp", str(cap.get("cp", ""))),
            ("Cpk", str(cap.get("cpk", ""))),
            (i18n.tr("report.ppm_total"), str(cap.get("ppm_total", ""))),
            (i18n.tr("report.grade"), str(cap.get("grade", ""))),
        ]
        for i, (k, v) in enumerate(rows):
            table.cell(i, 0).text = k
            table.cell(i, 1).text = v

    def _add_msa_results(self, result: Dict):
        doc = self.document
        vc = result.get("variance_components", {})
        pct = result.get("percentages", {})

        table = doc.add_table(rows=5, cols=2, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        rows = [
            (i18n.tr("report.ev"), str(vc.get("ev", ""))),
            (i18n.tr("report.av"), str(vc.get("av", ""))),
            ("GR&R", str(vc.get("grr", ""))),
            (i18n.tr("report.grr_pct"), f"{pct.get('grr_pct_tv', '')}%"),
            ("NDC", str(result.get("ndc", ""))),
        ]
        for i, (k, v) in enumerate(rows):
            table.cell(i, 0).text = k
            table.cell(i, 1).text = v

        acceptance = result.get("acceptance", {})
        doc.add_paragraph(f"{i18n.tr('report.judgment')}: {acceptance.get('message', '')} ({acceptance.get('level', '')})")

    def _add_hypothesis_results(self, result: Dict):
        doc = self.document
        table = doc.add_table(rows=4, cols=2, style="Light Grid Accent 1")
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        rows = [
            (i18n.tr("report.test_type"), result.get("test_type", "")),
            ("Statistic", str(result.get("t_statistic", result.get("f_statistic", "")))),
            ("P-value", str(result.get("p_value", ""))),
            (i18n.tr("report.conclusion"), i18n.tr("report.significant") if result.get("is_significant") else i18n.tr("report.not_significant")),
        ]
        for i, (k, v) in enumerate(rows):
            table.cell(i, 0).text = k
            table.cell(i, 1).text = v

    def _add_doe_results(self, result: Dict):
        doc = self.document
        effects = result.get("effects", result.get("main_effects", {}))

        p = doc.add_paragraph(i18n.tr("report.factor_effects_colon"))
        p.add_run("\n")  # Add newline after label
        run = p.add_run(
            ", ".join(f"{k}={v}" for k, v in list(effects.items())[:10])
        )
        run.font.size = Pt(10)

        doc.add_paragraph(f"R-squared = {result.get('r_squared', result.get('model_r_squared', ''))}")

    def _add_appendix(self):
        doc = self.document
        doc.add_heading(i18n.tr("report.appendix"), level=1)
        doc.add_paragraph(i18n.tr("report.appendix_formula"))
        doc.add_paragraph(i18n.tr("report.appendix_constant"))
        doc.add_paragraph(i18n.tr("report.appendix_note"))
