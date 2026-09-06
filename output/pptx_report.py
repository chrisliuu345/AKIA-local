"""
PowerPoint 报告生成器 — 基于 python-pptx 生成图文图表报告 (.pptx)

每分析一页：标题 + 图表（居中等比缩放），封面 + 页码页脚。
与 WordReport 共用 analyses 结构（analysis_type/params/result + charts PNG bytes）。
"""

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from config.settings import APP_NAME, APP_VERSION
from core.i18n import i18n
from ui.font_config import get_ui_font
from output.base_report_engine import BaseReportEngine

_SLIDE_W = 13.333
_SLIDE_H = 7.5
_GRAY = RGBColor(0x88, 0x88, 0x88)


class PptxReport(BaseReportEngine):
    """生成 PowerPoint (.pptx) 图文报告"""

    def __init__(self):
        super().__init__()
        self.font = get_ui_font()

    def generate(
        self,
        analyses: List[Dict[str, Any]],
        output_path: str,
        title: str = "",
        subtitle: str = "",
        charts: Optional[Dict[str, Any]] = None,
    ) -> str:
        if not title:
            title = i18n.tr("report.title")
        prs = Presentation()
        prs.slide_width = Inches(_SLIDE_W)
        prs.slide_height = Inches(_SLIDE_H)
        blank = prs.slide_layouts[6]

        self._add_cover(prs, blank, title, subtitle)
        total = len(analyses)
        for i, atype, params, result in self.iter_analyses(analyses):
            self._add_analysis_slide(
                prs, blank, atype, result, (charts or {}).get(atype), i + 1, total)

        prs.save(output_path)
        return output_path

    # ---------- 内部 ----------

    def _add_cover(self, prs, layout, title: str, subtitle: str):
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.2))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = title
        run.font.size = Pt(36)
        run.font.bold = True
        run.font.name = self.font
        if subtitle:
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            run2 = p2.add_run()
            run2.text = subtitle
            run2.font.size = Pt(18)
            run2.font.color.rgb = _GRAY
            run2.font.name = self.font

        box2 = slide.shapes.add_textbox(Inches(0.8), Inches(5.6), Inches(11.7), Inches(1.0))
        p3 = box2.text_frame.paragraphs[0]
        p3.alignment = PP_ALIGN.CENTER
        run3 = p3.add_run()
        run3.text = f"{APP_NAME} v{APP_VERSION}  ·  {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        run3.font.size = Pt(12)
        run3.font.color.rgb = _GRAY
        run3.font.name = self.font

    def _add_analysis_slide(self, prs, layout, atype: str, result: Dict,
                            chart_buf: Optional[Any], index: int, total: int):
        slide = prs.slides.add_slide(layout)

        title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(12.2), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = f"{index}、{atype}"
        run.font.size = Pt(26)
        run.font.bold = True
        run.font.name = self.font

        summary = (result or {}).get("summary")
        if isinstance(summary, str) and summary.strip():
            sbox = slide.shapes.add_textbox(Inches(0.55), Inches(1.0), Inches(12.2), Inches(0.55))
            sbox.text_frame.word_wrap = True
            sp = sbox.text_frame.paragraphs[0]
            srun = sp.add_run()
            srun.text = summary.strip().replace("\n", " · ")[:240]
            srun.font.size = Pt(11)
            srun.font.color.rgb = _GRAY
            srun.font.name = self.font

        img_top = Inches(1.7)
        img_box_w = Inches(12.3)
        img_box_h = Inches(5.0)
        if chart_buf is not None:
            data = BytesIO(chart_buf) if isinstance(chart_buf, bytes) else chart_buf
            data.seek(0)
            try:
                img = slide.shapes.add_picture(data, 0, 0)
                w_in = img.width / 914400
                h_in = img.height / 914400
                box_w, box_h = 12.3, 5.0
                scale = min(box_w / w_in, box_h / h_in)
                w_in, h_in = w_in * scale, h_in * scale
                img.left = Inches((_SLIDE_W - w_in) / 2)
                img.top = img_top
                img.width = Inches(w_in)
                img.height = Inches(h_in)
            except Exception:
                note = slide.shapes.add_textbox(
                    Inches(0.55), Inches(3.2), Inches(12.2), Inches(0.5))
                note.text_frame.paragraphs[0].add_run().text = "（图表数据不可用）"
        else:
            note = slide.shapes.add_textbox(Inches(0.55), Inches(3.4), Inches(12.2), Inches(0.5))
            note.text_frame.paragraphs[0].add_run().text = "（该分析无图表）"

        footer = slide.shapes.add_textbox(Inches(0.55), Inches(6.95), Inches(12.2), Inches(0.4))
        fp = footer.text_frame.paragraphs[0]
        frun = fp.add_run()
        frun.text = f"{index} / {total}    {APP_NAME} v{APP_VERSION}"
        frun.font.size = Pt(10)
        frun.font.color.rgb = _GRAY
