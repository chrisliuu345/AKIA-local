"""LPA 分层审核表引擎 — 模板生成 + Excel 导出"""

from typing import List, Dict
from ui.font_config import get_ui_font

LPA_TEMPLATES = {
    "5S": [
        ("整理", "不必要的物品是否已清除?", "目视"),
        ("整顿", "所有物品是否按规定位置摆放?", "目视"),
        ("清扫", "工作区域是否保持清洁?", "目视"),
        ("清洁", "清洁标准是否被执行?", "目视"),
        ("素养", "员工是否遵守5S规定?", "观察"),
    ],
    "安全": [
        ("PPE", "个人防护装备是否佩戴正确?", "目视"),
        ("消防", "灭火器是否在有效期内且无障碍?", "目视"),
        ("电气", "电线/插座是否完好无破损?", "目视"),
        ("标识", "安全标识是否清晰可见?", "目视"),
        ("应急", "紧急出口是否畅通?", "目视"),
    ],
    "质量": [
        ("首件检验", "首件是否经过检验并记录?", "记录审查"),
        ("过程参数", "关键参数是否在控制范围内?", "数据检查"),
        ("量具校准", "量具是否在有效校准期内?", "标签检查"),
        ("不合格品", "不合格品是否标识并隔离?", "目视"),
        ("追溯", "批次追溯信息是否完整?", "记录审查"),
    ],
    "标准化": [
        ("SOP", "作业指导书是否为最新版本?", "文件审查"),
        ("培训", "操作员是否经过岗位培训?", "记录审查"),
        ("点检", "设备点检是否按时完成?", "记录审查"),
        ("变更", "4M变更是否经过审批?", "文件审查"),
        ("记录", "生产记录是否完整准确?", "记录审查"),
    ],
}


class LPAEngine:
    @staticmethod
    def get_templates() -> Dict[str, list]:
        return LPA_TEMPLATES

    @staticmethod
    def generate_excel(
        processes: List[str],
        selected_categories: List[str] = None,
        auditors: List[str] = None,
        output_path: str = "LPA_Checklist.xlsx",
    ):
        """生成 LPA Excel 检查表"""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "LPA检查表"

        categories = selected_categories or list(LPA_TEMPLATES.keys())
        auditors = auditors or ["审核员A", "审核员B"]
        header_fill = PatternFill(
            start_color="1b5faa", end_color="1b5faa", fill_type="solid"
        )
        header_font = Font(name=get_ui_font(), color="FFFFFF", bold=True, size=10)
        pass_fill = PatternFill(
            start_color="c6efce", end_color="c6efce", fill_type="solid"
        )
        fail_fill = PatternFill(
            start_color="ffc7ce", end_color="ffc7ce", fill_type="solid"
        )
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        headers = ["工序", "类别", "检查项目", "检查方法"]
        for a in auditors:
            headers.append(f"{a}\n(OK/NG)")
        headers.append("备注")
        for j, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=j, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            cell.border = thin_border

        row = 2
        for proc in processes:
            for cat in categories:
                items = LPA_TEMPLATES.get(cat, [])
                for item_name, item_desc, item_method in items:
                    ws.cell(row=row, column=1, value=proc).border = thin_border
                    ws.cell(row=row, column=2, value=cat).border = thin_border
                    ws.cell(
                        row=row, column=3, value=f"{item_name}: {item_desc}"
                    ).border = thin_border
                    ws.cell(row=row, column=4, value=item_method).border = thin_border
                    for ai in range(len(auditors)):
                        cell = ws.cell(row=row, column=5 + ai, value="")
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal="center")
                    ws.cell(
                        row=row, column=5 + len(auditors), value=""
                    ).border = thin_border
                    row += 1

        widths = [12, 10, 40, 12] + [14] * len(auditors) + [20]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        audit_start_col = 5
        audit_end_col = 4 + len(auditors)
        from openpyxl.formatting.rule import CellIsRule

        ws.conditional_formatting.add(
            f"{get_column_letter(audit_start_col)}2:{get_column_letter(audit_end_col)}{row - 1}",
            CellIsRule(operator="equal", formula=['"OK"'], fill=pass_fill),
        )
        ws.conditional_formatting.add(
            f"{get_column_letter(audit_start_col)}2:{get_column_letter(audit_end_col)}{row - 1}",
            CellIsRule(operator="equal", formula=['"NG"'], fill=fail_fill),
        )

        wb.save(output_path)
        return output_path
